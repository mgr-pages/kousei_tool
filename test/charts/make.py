# グラフを含むデータファイルの原稿と、そのグラフを載せた成果物の模擬データを作る(文言・数はすべて架空)
#   python3 test/charts/make.py 出力フォルダ
#   chart.xlsx : シート「売れ行き」に表(月・台数)と棒グラフ(表題・軸の表題・データラベル)。グラフの値はこのシートのセル
#   chart.pptx : スライドの表題と、棒グラフ(表題・凡例・データラベル。値はグラフに保存されたもの)と円グラフ(項目名は凡例)
#   chart.docx : 本文と、段落の中の棒グラフ(PowerPoint の棒グラフと同じ部品)
#   final.pdf  : 成果物3ページ。グラフの文字をグラフのように並べたもの。違いは2か所：
#                1ページ目(Excel)の表題「月別の売れ行き」→「月別の売上」、2ページ目(PowerPoint)の棒グラフの値「150」→「160」
import sys, os, zipfile, re
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

out = sys.argv[1]; os.makedirs(out, exist_ok=True)
# ---- Excel ----
wb = Workbook(); ws = wb.active; ws.title = '売れ行き'
for r in [['月', '台数'], ['4月', 120], ['5月', 180], ['6月', 150]]: ws.append(r)
ch = BarChart(); ch.title = '月別の売れ行き（架空）'; ch.x_axis.title = '月'; ch.y_axis.title = '台数（台）'
ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True
ws.add_chart(ch, 'D2'); wb.save(f'{out}/_tmp.xlsx')
#   openpyxl はグラフに値の控え(キャッシュ)を書かないので、Excel と同じように書き足す
def cache_str(vals): return '<strCache><ptCount val="%d"/>%s</strCache>' % (len(vals), ''.join('<pt idx="%d"><v>%s</v></pt>' % (i, v) for i, v in enumerate(vals)))
def cache_num(vals): return '<numCache><formatCode>General</formatCode><ptCount val="%d"/>%s</numCache>' % (len(vals), ''.join('<pt idx="%d"><v>%s</v></pt>' % (i, v) for i, v in enumerate(vals)))
with zipfile.ZipFile(f'{out}/_tmp.xlsx') as zi, zipfile.ZipFile(f'{out}/chart.xlsx', 'w', zipfile.ZIP_DEFLATED) as zo:
    for n in zi.namelist():
        d = zi.read(n)
        if re.match(r'xl/charts/chart\d+\.xml$', n):
            x = d.decode()
            x = re.sub(r'(<tx><strRef><f>[^<]*</f>)', lambda m: m.group(1) + cache_str(['台数']), x, count=1)
            x = re.sub(r'(<cat><(?:num|str)Ref><f>[^<]*</f>)', lambda m: m.group(1) + cache_str(['4月', '5月', '6月']), x, count=1)
            x = re.sub(r'(<val><numRef><f>[^<]*</f>)', lambda m: m.group(1) + cache_num([120, 180, 150]), x, count=1)
            d = x.encode()
        zo.writestr(n, d)
os.remove(f'{out}/_tmp.xlsx')
# ---- PowerPoint ----
prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = '来店の状況'
cd = CategoryChartData(); cd.categories = ['東店', '西店', '南店']; cd.add_series('来店数', (120, 180, 150))
c = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(4.5), Inches(4), cd).chart
c.has_title = True; c.chart_title.text_frame.text = '店舗別の来店数（架空）'; c.has_legend = True; c.legend.position = XL_LEGEND_POSITION.BOTTOM; c.legend.include_in_layout = False
c.plots[0].has_data_labels = True; c.plots[0].data_labels.show_value = True
pd = CategoryChartData(); pd.categories = ['春', '夏', '秋']; pd.add_series('割合', (0.5, 0.3, 0.2))
p = s.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(5.2), Inches(1.5), Inches(4.3), Inches(4), pd).chart
p.has_title = True; p.chart_title.text_frame.text = '季節別の割合（架空）'; p.has_legend = True; p.legend.include_in_layout = False
prs.save(f'{out}/chart.pptx')
# ---- Word: python-docx はグラフを作れないので、PowerPoint の棒グラフの部品を段落の中に入れる ----
doc = Document(); doc.add_paragraph('来店数のご報告（架空）'); doc.add_paragraph('GRAPH_HERE'); doc.add_paragraph('以上です。'); doc.save(f'{out}/_tmp.docx')
with zipfile.ZipFile(f'{out}/chart.pptx') as zp: chart_xml = zp.read(sorted(n for n in zp.namelist() if re.match(r'ppt/charts/chart\d+\.xml$', n))[0])
inline = ('<w:r><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"><wp:extent cx="4114800" cy="3200400"/><wp:docPr id="10" name="グラフ 1"/>'
          '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
          '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdChart1"/>'
          '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')
with zipfile.ZipFile(f'{out}/_tmp.docx') as zi, zipfile.ZipFile(f'{out}/chart.docx', 'w', zipfile.ZIP_DEFLATED) as zo:
    for n in zi.namelist():
        d = zi.read(n)
        if n == 'word/document.xml': d = re.sub(rb'<w:r>(?:(?!<w:r>).)*?GRAPH_HERE</w:t></w:r>', inline.encode(), d, flags=re.S)
        if n == 'word/_rels/document.xml.rels': d = d.replace(b'</Relationships>', b'<Relationship Id="rIdChart1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="charts/chart1.xml"/></Relationships>')
        if n == '[Content_Types].xml': d = d.replace(b'</Types>', b'<Override PartName="/word/charts/chart1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/></Types>')
        zo.writestr(n, d)
    zo.writestr('word/charts/chart1.xml', chart_xml)
os.remove(f'{out}/_tmp.docx')
# ---- 成果物: グラフの文字をグラフのように並べる(A4横・3ページ) ----
pdfmetrics.registerFont(TTFont('ipag', '/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf'))
cv = canvas.Canvas(f'{out}/final.pdf', pagesize=(842, 595))
def t(x, y, s, size=12): cv.setFont('ipag', size); cv.drawString(x, y, s)
def bars(x0, y0, vals, labels, w=50, gap=40):
    for k, (v, lab) in enumerate(zip(vals, labels)):
        x = x0 + k * (w + gap); h = v * 1.2; cv.setFillGray(0.6); cv.rect(x, y0, w, h, fill=1, stroke=0); cv.setFillGray(0.1); t(x + 12, y0 + h + 6, str(v), 11); t(x + 10, y0 - 18, lab, 11)
cv.setFillGray(0.1)
t(60, 520, '月別の売上（架空）', 18); t(30, 330, '台数（台）', 11); bars(110, 180, [120, 180, 150], ['4月', '5月', '6月']); t(230, 130, '月', 11); cv.showPage()   # Excel のグラフ(表題が違う)
t(60, 530, '来店の状況', 22)
t(60, 490, '店舗別の来店数（架空）', 14); bars(80, 180, [120, 180, 160], ['東店', '西店', '南店']); t(180, 130, '来店数', 11)   # PowerPoint の棒グラフ(値が違う)
t(500, 490, '季節別の割合（架空）', 14); cv.setFillGray(0.6); cv.circle(600, 330, 90, fill=1, stroke=0); cv.setFillGray(0.1); t(540, 200, '春　夏　秋', 11); cv.showPage()
t(60, 530, '来店数のご報告（架空）', 14); t(60, 470, '店舗別の来店数（架空）', 13); bars(80, 200, [120, 180, 150], ['東店', '西店', '南店']); t(180, 150, '来店数', 11); t(60, 90, '以上です。', 12); cv.showPage()
cv.save()
print('ok', out)
