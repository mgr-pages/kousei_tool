# データファイルの原稿(Excel・Word・PowerPoint・テキスト・CSV)と、それを組んだ成果物PDFの模擬データを作る(自動テスト用。文言はすべて架空)
#   python3 test/office/make.py 出力フォルダ
#   → 1_price.xlsx・2_notice.docx・3_set.pptx・4_contact.txt・5_shops.csv(原稿) と final.pdf(成果物 A4縦 3ページ)
#   成果物には、決まった違いを入れてある(正解は test/office/gt.json)。原稿の非表示の行・列・シート・スライド、Wordの削除の変更履歴は成果物に載せない
import sys, os, datetime
from openpyxl import Workbook
from openpyxl.comments import Comment
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Pt, Emu

out = sys.argv[1] if len(sys.argv) > 1 else 'testdata/office'
os.makedirs(out, exist_ok=True)
FONT = '/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf'

# ---- Excel: 価格表(見出しの結合セル・表示形式・非表示の行と列・コメント・非表示のシート) ----
ROWS = [
    ('B-101', 'やさしい手帳 A5', '120枚', 1280, '定番'),
    ('B-102', 'やさしい手帳 B6', '96枚', 1050, ''),
    ('B-203', '方眼ノート 5冊組', '5冊', 980, '人気'),
    ('B-204', '色えんぴつ 12色', '12本', 1650, ''),
    ('B-305', '消しゴム 角型', '1個', 120, ''),
    ('B-306', '消しゴム 旧型', '1個', 100, '廃番'),      # 非表示の行
    ('B-410', 'クリップ 小箱', '50個', 220, ''),
    ('B-411', 'ふせん 中', '4色', 360, ''),
    ('B-520', '定規 15cm', '1本', 180, ''),
    ('B-521', '定規 30cm', '1本', 260, ''),
    ('B-630', 'ペンケース 布', '1個', 1480, '新色'),
    ('B-631', 'ペンケース 革風', '1個', 2200, ''),
]
wb = Workbook()
ws = wb.active
ws.title = '価格表'
ws['A1'] = '2027年度　文具価格表'
ws.merge_cells('A1:E1')
ws.append(['品番', '商品名', '容量', '価格', '備考', '原価'])
for i, (a, b, c, d, e) in enumerate(ROWS):
    ws.append([a, b, c, d, e or None, round(d * 0.6)])
    ws.cell(row=ws.max_row, column=4).number_format = '#,##0"円"'
    if a == 'B-306':
        ws.row_dimensions[ws.max_row].hidden = True
    if a == 'B-204':
        ws.cell(row=ws.max_row, column=4).comment = Comment('価格は改定後の値', '担当')
ws.column_dimensions['F'].hidden = True                     # 非表示の列(原価)
ws.append(['発売日', datetime.datetime(2027, 4, 1)])
ws.cell(row=ws.max_row, column=2).number_format = 'yyyy"年"m"月"d"日"'
ws.append(['割引率', 0.15])
ws.cell(row=ws.max_row, column=2).number_format = '0%'
memo = wb.create_sheet('メモ')
memo['A1'] = '社外秘：仕入れ先との打ち合わせメモ'
memo.sheet_state = 'hidden'                                  # 非表示のシート
wb.save(os.path.join(out, '1_price.xlsx'))

# ---- Word: お知らせ(見出しへのコメント・変更履歴・表) ----
doc = Document()
h = doc.add_paragraph()
hr = h.add_run('春の文具フェアのお知らせ')
doc.add_comment([hr], text='見出しは大きく', author='担当', initials='T')
doc.add_paragraph('日ごろより文具ひろばをご利用いただき、ありがとうございます。今年も春の文具フェアを開催します。')
p = doc.add_paragraph()
p.add_run('会場は')
def change(tag, text, rid):   # 変更履歴(w:ins / w:del)
    el = OxmlElement('w:' + tag)
    el.set(qn('w:id'), str(rid)); el.set(qn('w:author'), '担当'); el.set(qn('w:date'), '2027-01-10T09:00:00Z')
    r = OxmlElement('w:r'); t = OxmlElement('w:delText' if tag == 'del' else 'w:t'); t.text = text
    r.append(t); el.append(r)
    return el
p._p.append(change('del', '旧館', 101))
p._p.append(change('ins', '新館', 102))
p.add_run('の2階です。ご来場の際は、正面入口をご利用ください。')
doc.add_paragraph('期間中は、対象の商品がすべて特別価格になります。数に限りがありますので、お早めにお越しください。')
t = doc.add_table(rows=3, cols=2)
for r, (k, v) in enumerate([('日時', '2027年3月20日（土）〜3月22日（月）'), ('会場', '文具ひろば 新館2階'), ('定員', '各回先着50名')]):
    t.cell(r, 0).text = k; t.cell(r, 1).text = v
doc.add_paragraph('皆さまのご来場をお待ちしております。')
doc.save(os.path.join(out, '2_notice.docx'))

# ---- PowerPoint: セットの案内(表・非表示のスライド) ----
prs = Presentation()
s1 = prs.slides.add_slide(prs.slide_layouts[5])   # タイトルのみ
s1.shapes.title.text = 'フェア限定セット'
tb = s1.shapes.add_textbox(Emu(457200), Emu(1600000), Emu(8000000), Emu(600000))
tb.text_frame.text = '3点まとめて15%お得'
tbl = s1.shapes.add_table(3, 2, Emu(457200), Emu(2400000), Emu(8000000), Emu(1200000)).table
for r, row in enumerate([('セット名', '内容'), ('入学準備セット', '手帳・ノート・ペンケース'), ('事務用セット', 'クリップ・ふせん・定規')]):
    for c, v in enumerate(row):
        tbl.cell(r, c).text = v
s2 = prs.slides.add_slide(prs.slide_layouts[5])
s2.shapes.title.text = '社内用：在庫の確認を先に'
s2._element.set('show', '0')                       # 非表示のスライド
prs.save(os.path.join(out, '3_set.pptx'))

# ---- テキスト(Shift_JIS)・CSV(Shift_JIS、カンマを含むセル) ----
with open(os.path.join(out, '4_contact.txt'), 'wb') as f:
    f.write('お問い合わせ先\r\n文具ひろば フェア事務局\r\n電話 000-0000-0000（平日10時〜17時）\r\nメール info@example.com\r\n'.encode('cp932'))
with open(os.path.join(out, '5_shops.csv'), 'wb') as f:
    f.write('店舗名,所在地,電話\r\n本店,"架空県みどり市1-2-3, みどりビル1階",000-111-2222\r\n駅前店,架空県みどり市4-5-6,000-333-4444\r\n'.encode('cp932'))

# ---- 成果物: A4縦 3ページ(配置・改行・字の大きさは原稿と違う) ----
#   reportlab で作る(PyMuPDF で字を入れると、文字データの「-」が字の区切りの記号(U+00AD)になり、読み取りで消えるため)
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
pdfmetrics.registerFont(TTFont('ipag', FONT))
H = 842
class Page:
    def __init__(self, c): self.c = c
def new_page():
    return Page(cv)
def box(page, rect, text, size, color=(0.15, 0.15, 0.15), align=0):   # rect: 左上を原点とした (x0, y0, x1, y1)。幅で折り返す
    c = page.c; x0, y0, x1, y1 = rect
    c.setFont('ipag', size); c.setFillColorRGB(*color)
    lines, cur = [], ''
    for ch in text:
        if cur and pdfmetrics.stringWidth(cur + ch, 'ipag', size) > x1 - x0: lines.append(cur); cur = ''
        cur += ch
    lines.append(cur)
    assert y0 + len(lines) * size * 1.3 <= y1 + size * 0.3, (text, rect)
    for k, ln in enumerate(lines):
        c.drawString(x0, H - (y0 + size + k * size * 1.3), ln)
def table(page, x, y, widths, rows, size=9, rh=18):
    c = page.c
    for row in rows:
        cx = x
        for w, v in zip(widths, row):
            c.setStrokeColorRGB(0.6, 0.6, 0.6); c.setLineWidth(0.5); c.rect(cx, H - (y + rh), w, rh)
            if v:
                box(page, (cx + 3, y + 3, cx + w - 2, y + rh), v, size)
            cx += w
        y += rh
    return y
cv = canvas.Canvas(os.path.join(out, 'final.pdf'), pagesize=(595, H))
W = [55, 150, 60, 70, 60]
HEAD = ('品番', '商品名', '容量', '価格', '備考')
# 1ページ目: お知らせ(Word) と 価格表の前半(Excel)
p = new_page()
box(p, (40, 40, 555, 75), '春の文具フェアのお知らせ', 22, (0.1, 0.3, 0.6))
box(p, (40, 90, 380, 140), '日ごろより文具ひろばをご利用いただき、ありがとうございます。今年も春の文具フェアを開催します。', 10.5)
box(p, (40, 145, 380, 195), '会場は新館の2階です。ご来店の際は、正面入口をご利用ください。', 10.5)                 # ご来場→ご来店
box(p, (40, 200, 380, 240), '期間中は、対象の商品がすべて特別価格になります。', 10.5)                              # 2文目が抜けている
y = table(p, 40, 250, [60, 260], [('日時', '2027年3月20日（土）〜3月22日（月）'), ('会場', '文具ひろば 新館2階'), ('定員', '各回先着５０名')], 10, 20)   # 50→５０(全角)
box(p, (40, y + 10, 555, y + 30), '皆さまのご来場をお待ちしております。', 10.5)
box(p, (40, 420, 555, 440), '2027年度　文具価格表', 14, (0.1, 0.3, 0.6))
rows1 = [HEAD,
         ('B-101', 'やさしい手帳 A5', '120枚', '1,380円', '定番'),     # 1,280円→1,380円
         ('B-102', 'やさしい手帳 B6', '96枚', '1050円', ''),          # 桁区切りなし(違いとしない)
         ('B-203', '方眼ノート 5冊組', '5冊', '980円', '人気'),
         ('B-204', '色えんぴつ 12色', '１２本', '1,650円', ''),       # 12本→１２本(全角)
         ('B-305', '消しゴム 角型', '1個', '120円', '')]
table(p, 40, 450, W, rows1)
# 2ページ目: 価格表の後半(見出しをくり返す) と セットの案内(PowerPoint)
cv.showPage(); p = new_page()
box(p, (40, 40, 555, 60), '2027年度　文具価格表（つづき）', 14, (0.1, 0.3, 0.6))
rows2 = [HEAD,
         # B-410 クリップ 小箱 が抜けている
         ('B-411', 'ふせん 中', '4色', '360円', ''),
         ('B-412', 'ふせん 特大', '6色', '2,000円', ''),              # 原稿にない行
         ('B-520', '定規 15cm', '1本', '180円', ''),
         ('B-521', '定規 30cm', '1本', '260円', ''),
         ('B-630', 'ペンケース 布', '1個', '1,480円', '新色'),
         ('B-631', 'ペンケース 革風', '1個', '2,200円', '')]
y = table(p, 40, 70, W, rows2)
y = table(p, 40, y + 12, [80, 150], [('発売日', '2027年4月1日'), ('割引率', '15%')])
box(p, (40, y + 8, 555, y + 24), '※価格はすべて税込です。', 9)                                        # 原稿にない注記
box(p, (40, 420, 555, 450), 'フェア限定セット', 20, (0.6, 0.2, 0.1))
box(p, (40, 455, 555, 475), '3点まとめて20%お得', 12)                                                # 15%→20%
table(p, 40, 485, [140, 260], [('セット名', '内容'), ('入学準備セット', '手帳・ノート・ペンケース'), ('事務用セット', 'クリップ・ふせん・定規')], 10, 20)
# 3ページ目: 問い合わせ(テキスト) と 店舗一覧(CSV)
cv.showPage(); p = new_page()
box(p, (40, 40, 555, 60), 'お問い合わせ先', 14, (0.1, 0.3, 0.6))
box(p, (40, 66, 555, 82), '文具ひろば フェア事務局', 11)
box(p, (40, 86, 555, 102), '電話 000-0000-0000（平日10時〜18時）', 11)                                # 17時→18時
box(p, (40, 106, 555, 122), 'メール info@example.com', 11)
table(p, 40, 150, [70, 250, 110], [('店舗名', '所在地', '電話'), ('本店', '架空県みどり市1-2-3, みどりビル1階', '000-111-2222'), ('駅前店', '架空県みどり市4-5-6', '000-333-4445')], 9.5, 20)   # 4444→4445
cv.showPage(); cv.save()
print('ok', out)
